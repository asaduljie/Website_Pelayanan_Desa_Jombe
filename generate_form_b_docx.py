import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_form_b():
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base font
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    def set_cell_background(cell, fill_hex):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        tcPr.append(shd)

    def add_section_header(title):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, 'D1E7DD')  # Soft sage green header
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(15, 23, 42)
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_item_box(label, content_paragraphs):
        # Label paragraph
        p_lbl = doc.add_paragraph()
        p_lbl.paragraph_format.space_before = Pt(6)
        p_lbl.paragraph_format.space_after = Pt(3)
        run_lbl = p_lbl.add_run(label)
        run_lbl.bold = True
        run_lbl.italic = True
        run_lbl.font.size = Pt(11)
        run_lbl.font.color.rgb = RGBColor(15, 23, 42)

        # Content table box
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.rows[0].cells[0]
        cell.width = Inches(6.5)
        set_cell_background(cell, 'FFFFFF')

        # Set thin borders for the box
        tcPr = cell._element.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>
                <w:left w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="94A3B8"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)

        # Add contents
        for idx, text in enumerate(content_paragraphs):
            if idx == 0:
                p = cell.paragraphs[0]
            else:
                p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(text)
            run.font.size = Pt(10.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ═════════════════════ COVER / HEADER PAGE ═════════════════════
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_h1 = p_title.add_run("FORM B\n")
    r_h1.bold = True
    r_h1.font.size = Pt(28)
    r_h1.font.color.rgb = RGBColor(30, 41, 59)
    
    r_sub = p_title.add_run("Program Kerja\n")
    r_sub.italic = True
    r_sub.font.size = Pt(16)
    r_sub.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    p_prog = doc.add_paragraph()
    p_prog.paragraph_format.line_spacing = 1.2
    r_num = p_prog.add_run("Program kerja ke- [04]\n")
    r_num.font.size = Pt(14)
    r_num.bold = False

    r_prog_title = p_prog.add_run("[LENTERA DESA (LAYANAN ELEKTRONIK TERINTEGRASI ADMINISTRASI DESA)]\n")
    r_prog_title.bold = True
    r_prog_title.font.size = Pt(16)
    r_prog_title.font.color.rgb = RGBColor(15, 23, 42)

    # Horizontal rule
    p_hr = doc.add_paragraph()
    p_hr.paragraph_format.space_after = Pt(14)
    p_hr_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="0F172A"/></w:pBdr>')
    p_hr._p.get_or_add_pPr().append(p_hr_border)

    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.line_spacing = 1.3
    p_meta.add_run("Kabupaten        : Jeneponto\n")
    p_meta.add_run("Kecamatan        : Turatea\n")
    p_meta.add_run("Desa             : Jombe\n")
    p_meta.add_run("Dosen Pembimbing : Muhammad Ridha, S.H., M.A.\n")

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    # Note about Screenshot
    p_img_note = doc.add_paragraph()
    p_img_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_note = p_img_note.add_run("[ Tangkapan Layar Tampilan Portal Utama LENTERA DESA (https://www.lenteradesajombe.biz.id) ]")
    r_note.italic = True
    r_note.font.size = Pt(9.5)
    r_note.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # ═════════════════════ PAGE 2: BAGIAN I & II ═════════════════════
    p_h2 = doc.add_paragraph()
    p_h2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_h2 = p_h2.add_run("FORM B\n")
    r_h2.bold = True
    r_h2.font.size = Pt(16)
    r_h2_sub = p_h2.add_run("PELAKSANAAN, HASIL, DAN DAMPAK PROGRAM KERJA KKN")
    r_h2_sub.font.size = Pt(11)
    r_h2_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_line = doc.add_paragraph()
    p_line_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="334155"/></w:pBdr>')
    p_line._p.get_or_add_pPr().append(p_line_border)

    add_section_header("BAGIAN I: IDENTITAS PROGRAM KERJA")

    add_item_box("B1. Nama Posko KKN", [
        "KKN 79 UIN Alauddin Makassar Posko 6 Desa Jombe"
    ])

    add_item_box("B2. Nama Program Kerja", [
        "LENTERA DESA (Layanan Elektronik Terintegrasi Administrasi Desa)"
    ])

    add_item_box("B3. Kategori Program", [
        "☑ Lainnya: Digitalisasi Pelayanan Publik / Administrasi Pemerintahan Desa (E-Government / Smart Village)"
    ])

    add_item_box("B4. Status Program", [
        "☑ Program Unggulan"
    ])

    add_section_header("BAGIAN II: DESKRIPSI PELAKSANAAN PROGRAM")

    add_item_box("B5. Latar Belakang Singkat Program", [
        "Program ini diinisiasi untuk menjawab tantangan tata kelola administrasi kependudukan di Desa Jombe yang selama ini masih bersifat konvensional. Warga harus meluangkan waktu dan biaya perjalanan menuju kantor desa hanya untuk mencari informasi syarat surat, mengantre, ataupun memeriksa apakah surat sudah selesai ditandatangani. Di sisi lain, arsip permohonan surat masih bertumpu pada buku register manual yang berisiko tercecer atau rusak.",
        "Oleh karena itu, dikembangkan inovasi sistem terintegrasi 'LENTERA DESA' (Website Pelayanan Mandiri dan Aplikasi Android) agar masyarakat Desa Jombe dapat mengajukan permohonan surat administrasi secara langsung secara mandiri, melacak status berkas secara transparan, serta mengunduh dokumen surat resmi berformat PDF yang dilengkapi tanda tangan elektronik (TTE) berbasis QR Code tanpa terkendala jarak dan waktu."
    ])

    # ═════════════════════ PAGE 3: BAGIAN III ═════════════════════
    add_item_box("B6. Tujuan Program Kerja", [
        "1. Menyediakan platform pelayanan administrasi publik digital yang memudahkan warga mengajukan surat keterangan mandiri secara daring (online) 24 jam.",
        "2. Meningkatkan transparansi dan akuntabilitas tata kelola birokrasi desa melalui fitur pelacakan status berkas permohonan secara real-time.",
        "3. Memodernisasi tata kelola administrasi kantor desa dengan dashboard operator untuk verifikasi berkas digital dan penerbitan surat resmi otomatis ber-kop Desa Jombe.",
        "4. Menjamin keabsahan dokumen kependudukan melalui sistem verifikasi Tanda Tangan Elektronik (TTE) berbasis QR Code / Barcode yang dapat divalidasi secara publik.",
        "5. Menyediakan kanal pengaduan dan aspirasi warga serta portal profil desa terintegrasi berbasis data resmi BPS Kecamatan Turatea 2025."
    ])

    add_item_box("B7. Sasaran Program", [
        "Masyarakat Desa Jombe (lintas 5 dusun: Dusun Jombe Utara, Dusun Jombe Tengah, Dusun Jombe Selatan, Dusun Tompo Balang, dan Dusun Muncu-muncu) sebagai pengguna layanan mandiri, serta aparatur Pemerintah Desa Jombe (Sekretaris Desa dan Operator/Admin Desa) sebagai pengelola dan verifikator sistem."
    ])

    add_section_header("BAGIAN III: WAKTU, LOKASI METODE DAN MITRA PROGRAM")

    add_item_box("B8. Waktu dan Lokasi Pelaksanaan", [
        "• Tahap Demo/Pengenalan Awal: Kamis, 20 Agustus 2026 (pada saat Seminar Program Kerja KKN, berupa demonstrasi alur pengajuan surat online, sistem pelacakan mandiri warga, dan dashboard operator desa).",
        "• Tahap Pelatihan & Sosialisasi Utama: Jumat, 18 September 2026 pukul 08.30 WITA (dirangkaikan secara terintegrasi dengan Seminar Hasil Program Kerja KKN, mencakup bimbingan teknis bagi operator desa dan sosialisasi kepada perwakilan warga).",
        "• Lokasi Pelaksanaan: Kantor Desa Jombe, Kecamatan Turatea, Kabupaten Jeneponto."
    ])

    add_item_box("B9. Metode Pelaksanaan Program", [
        "1. Analisis Kebutuhan & Regulasi: Identifikasi jenis surat kependudukan yang paling sering dibutuhkan warga (SKU, Domisili, SKTM, Keterangan Kelahiran/Kematian, Pengantar SKCK) beserta format blangko resmi Desa Jombe.",
        "2. Pengembangan Perangkat Lunak (Software Engineering): Perancangan sistem berbasis web modern (Next.js 14, TypeScript, Tailwind CSS) dan Aplikasi Mobile Android (Flutter) dengan arsitektur backend API terenkripsi dan basis data relasional.",
        "3. Pengujian Sistem (Quality Assurance): Pengujian alur registrasi warga, pengunggahan berkas digital (KTP/KK), verifikasi operator, dan penerbitan dokumen PDF ber-QR Code.",
        "4. Bimbingan Teknis (Bimtek) & Pendampingan Operator: Pelatihan operasional manajemen permohonan, verifikasi syarat, hingga penerbitan surat kepada aparatur kantor desa.",
        "5. Sosialisasi & Edukasi Masyarakat: Penjelasan alur layanan mandiri digital kepada para Kepala Dusun, kader posyandu, dan tokoh masyarakat."
    ])

    add_item_box("B10. Mitra Program yang Terlibat", [
        "• Nama Mitra: Pemerintah Desa Jombe (Kepala Desa JUSMAEDY, S.Pd, Sekretaris Desa SYAMSUL RISWAN, jajaran Kepala Seksi, Kepala Urusan, dan 5 Kepala Dusun).",
        "• Peran Mitra: Menyediakan data pelayanan publik dan format surat baku desa, mendampingi perancangan alur verifikasi berkas, memfasilitasi sarana pelatihan di kantor desa, serta bertindak sebagai operator/administrator pengelola resmi sistem LENTERA DESA pasca-KKN."
    ])

    # ═════════════════════ PAGE 4: BAGIAN IV & V ═════════════════════
    add_section_header("BAGIAN IV: OUTPUT PROGRAM")

    add_item_box("B11. Rangkaian Kegiatan yang Dilaksanakan", [
        "1. Pengumpulan format surat dinas resmi dan inventarisasi alur pelayanan administrasi kependudukan di Kantor Desa Jombe.",
        "2. Perancangan dan coding arsitektur sistem Website Pelayanan Publik (Next.js 14), Aplikasi Android (Flutter), dan Backend REST API.",
        "3. Implementasi sistem keamanan data warga (enkripsi AES, verifikasi signature file, otorisasi multi-role warga & operator).",
        "4. Integrasi fitur Tracking Status Surat Real-time, Verifikasi Keaslian Dokumen ber-QR Code, dan Formulir Pengaduan Masyarakat.",
        "5. Pengujian fungsi sistem (alpha/beta testing) dan simulasi penerbitan surat keterangan secara digital.",
        "6. Pelatihan intensif pengoperasian sistem kepada Sekretaris Desa dan staf operator administrasi desa.",
        "7. Pelaksanaan sosialisasi alur layanan publik digital LENTERA DESA kepada para Kepala Dusun, kader, dan perwakilan masyarakat."
    ])

    add_item_box("B12. Output Program yang Dihasilkan", [
        "1. Tersedianya Website Resmi Pelayanan Mandiri Desa Jombe (https://www.lenteradesajombe.biz.id) yang aktif dan dapat diakses 24 jam.",
        "2. Tersedianya Aplikasi Mobile Android (APK Flutter) yang terhubung langsung dengan basis data pelayanan desa.",
        "3. Tersedianya Dashboard Operator Desa untuk mengelola, memverifikasi berkas, dan menerbitkan surat keterangan berformat PDF resmi.",
        "4. Tersedianya Modul Verifikasi Keaslian Tanda Tangan Elektronik (TTE) berbasis QR Code untuk mencegah pemalsuan surat desa.",
        "5. Tersedianya Kanal Pengaduan Aspirasi Warga dan Portal Profil Interaktif Desa Jombe berbasis data resmi BPS Kecamatan Turatea 2025.",
        "6. Terlatihnya aparatur Pemerintah Desa Jombe dalam mengoperasikan dan mengelola sistem pelayanan digital secara mandiri."
    ])

    add_item_box("B13. Jumlah dan Partisipasi Peserta", [
        "• Pemerintah Desa yang Dilatih: Sekretaris Desa dan 1 staf operator administrasi Kantor Desa Jombe (pendampingan teknis intensif).",
        "• Peserta Sosialisasi & Edukasi: Dihadiri secara aktif oleh sekitar 33 orang perwakilan target yang terdiri atas para Kepala Dusun (Dusun Jombe Utara, Tengah, Selatan, Tompo Balang, Muncu-muncu), para kader posyandu/desa, tokoh pemuda, dan masyarakat umum."
    ])

    add_section_header("BAGIAN V: OUTCOME DAN DAMPAK AWAL")

    add_item_box("B14. Perubahan yang Teramati pada Sasaran", [
        "• Efisiensi Birokrasi Desa: Aparatur pemerintah desa kini memiliki platform digital terpusat untuk memproses permohonan surat kependudukan secara cepat tanpa harus mengetik ulang blangko manual satu per satu.",
        "• Kemudahan Akses Masyarakat: Warga tidak lagi terbebani harus bolak-balik ke kantor desa hanya untuk menanyakan kelengkapan syarat atau mengecek apakah surat sudah jadi. Warga dapat memantau progres permohonan langsung dari gawai (smartphone) mereka.",
        "• Penguatan Peran Kepala Dusun: Para kepala dusun dan kader desa memahami alur layanan digital sehingga mampu menjadi jembatan informasi dan membimbing warga di dusun masing-masing yang membutuhkan bantuan pengajuan surat.",
        "• Transparansi & Keamanan Arsip: Seluruh rekam jejak permohonan tersimpan rapi secara digital dan permanen, memudahkan pencatatan register surat masuk/keluar di kantor desa."
    ])

    # ═════════════════════ PAGE 5: BAGIAN VI & VII ═════════════════════
    add_item_box("B15. Respon dan Testimoni Masyarakat", [
        "Pemerintah Desa Jombe menyambut dengan sangat antusias kehadiran sistem LENTERA DESA karena sejalan dengan visi modernisasi desa berbasis swasembada. Kepala Desa dan Sekretaris Desa menyampaikan apresiasi tinggi karena sistem ini secara nyata memangkas beban kerja administrasi berulang di kantor desa.",
        "Masyarakat dan para kepala dusun merespons sangat positif karena sistem dinilai sangat praktis, hemat waktu, dan transparan, terutama bagi warga yang bermukim di dusun yang jauh dari kantor desa."
    ])

    add_section_header("BAGIAN VI: KEBERLANJUTAN PROGRAM")

    add_item_box("B16. Potensi Keberlanjutan Program", [
        "Potensi keberlanjutan program sangat tinggi. Sistem LENTERA DESA telah di-deploy pada infrastruktur server cloud yang stabil dan terjangkau, menggunakan basis data cloud terpusat, serta seluruh hak akses, panduan operasional, dan akun administrator telah diserahterimakan secara penuh kepada Pemerintah Desa Jombe.",
        "Sistem ini dirancang modular sehingga mudah dikembangkan lebih lanjut oleh pemerintah desa sesuai dengan penambahan kebutuhan jenis layanan surat di masa depan."
    ])

    add_item_box("B17. Pihak yang Dapat Melanjutkan Program", [
        "Aparatur Pemerintah Desa Jombe (khususnya Sekretaris Desa sebagai penanggung jawab administrasi, dan Staf Operator/Admin IT Desa sebagai pelaksana harian)."
    ])

    add_section_header("BAGIAN VII: REFLEKSI DAN PEMBELAJARAN")

    add_item_box("B18. Tantangan dalam Pelaksanaan Program", [
        "1. Keterbatasan sumber daya manusia aparatur yang standby secara purna waktu di kantor desa (terpusat pada Sekretaris Desa dan 1 staf administrasi).",
        "2. Variasi tingkat literasi digital sebagian masyarakat di beberapa dusun, khususnya warga lanjut usia.",
        "3. Tantangan sinkronisasi agenda sosialisasi agar dapat menghadirkan seluruh perwakilan dari kelima wilayah dusun secara efektif dan efisien."
    ])

    add_item_box("B19. Solusi yang Dilakukan", [
        "1. Melakukan bimbingan teknis secara intensif dan pendampingan tatap muka (one-on-one coaching) kepada operator desa hingga mahir mengelola seluruh alur verifikasi berkas dan penerbitan surat.",
        "2. Merancang antarmuka (UI/UX) website dan aplikasi mobile dengan prinsip kesederhanaan, ukuran font yang jelas, bahasa Indonesia yang ramah, serta alur pengajuan yang ringkas dan mudah dipahami oleh masyarakat awam.",
        "3. Menyediakan alur alternatif (pendampingan operator/Plan B) di kantor desa bagi warga lansia yang membutuhkan bantuan penginputan berkas secara langsung.",
        "4. Mengintegrasikan momentum Sosialisasi Utama LENTERA DESA bersamaan dengan agenda Seminar Hasil KKN sehingga dihadiri langsung oleh seluruh Kepala Dusun, kader posyandu, dan tokoh masyarakat dari 5 dusun."
    ])

    # ═════════════════════ PAGE 6: BAGIAN VIII ═════════════════════
    add_item_box("B20. Pembelajaran Penting dari Program", [
        "Implementasi transformasi digital di tingkat desa tidak hanya bertumpu pada kecanggihan teknologi perangkat lunak, tetapi menuntut kesiapan kapasitas sumber daya manusia aparatur dan kemudahan adopsi bagi masyarakat. Pendekatan humanis, penyediaan panduan visual yang jelas, serta kolaborasi erat bersama kepala dusun sebagai garda terdepan kewilayahan merupakan faktor kunci keberhasilan penerapan e-Government di tingkat desa."
    ])

    add_section_header("BAGIAN VIII: DOKUMENTASI PROGRAM")

    add_item_box("B21. Dokumentasi Kegiatan", [
        "☐ Foto 1: Pendemoan fitur Website Pelayanan Mandiri LENTERA DESA, alur pengajuan surat online, dan sistem lacak berkas saat Seminar Program Kerja KKN (20 Agustus 2026).",
        "☐ Foto 2: Proses pendampingan dan bimbingan teknis (Bimtek) operasional Dashboard Operator Desa bersama Sekretaris Desa dan staf kantor desa.",
        "☐ Foto 3: Pelaksanaan Sosialisasi LENTERA DESA yang dirangkaikan dengan Seminar Hasil KKN bersama Kepala Desa, Sekretaris Desa, seluruh Kepala Dusun, dan Kader Posyandu se-Desa Jombe (18 September 2026).",
        "☐ Foto 4: Contoh hasil cetak Surat Keterangan resmi Desa Jombe berformat PDF yang dilengkapi Tanda Tangan Elektronik (TTE) dan QR Code validasi keaslian dokumen."
    ])

    # Footer note
    p_footer_rule = doc.add_paragraph()
    p_footer_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="94A3B8"/></w:pBdr>')
    p_footer_rule._p.get_or_add_pPr().append(p_footer_border)

    p_fn = doc.add_paragraph()
    p_fn.paragraph_format.space_before = Pt(6)
    p_fn.add_run("CATATAN WAJIB BAGI MAHASISWA:\n").bold = True
    p_fn.add_run("• Satu FORM B = satu program kerja.\n")
    p_fn.add_run("• Diisi setelah program selesai, bukan di akhir KKN.\n")
    p_fn.add_run("• Jawaban bersifat faktual, bukan opini.\n")
    p_fn.add_run("• FORM B inilah yang langsung menjadi BAB Program Kerja di buku laporan KKN.\n")
    p_fn.add_run("• FORM B program unggulan menjadi bahan utama artikel jurnal pengabdian masyarakat.\n")

    output_path = "c:\\Users\\Lenovo\\Music\\kkn desa\\FORM_B_PROGRAM_KERJA_LENTERA_DESA_REVISI.docx"
    doc.save(output_path)
    print("SUCCESS: File saved to", output_path)

if __name__ == "__main__":
    create_form_b()
